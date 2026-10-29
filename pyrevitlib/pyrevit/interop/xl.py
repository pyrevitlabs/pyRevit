"""Read and write Excel files on both Python engines.

CPython reads and writes ``.xlsx`` workbooks with openpyxl and reads legacy
``.xls`` workbooks with the vendored xlrd, choosing the reader from the file
signature and falling back to the extension. IronPython keeps xlrd and
xlsxwriter for every workbook.

Invariant:
    ``load`` returns the same values on both engines: openpyxl cells are
    normalized to what xlrd reports (blank as ``''``, numbers and dates as
    ``float`` Excel serials, booleans as ``int``, errors as xlrd error codes).
"""

import datetime
import os.path as op

from pyrevit.compat import IRONPY


if IRONPY:
    import xlsxwriter
else:
    import openpyxl


_ZIP_SIGNATURE = b"PK\x03\x04"
_OLE2_SIGNATURE = b"\xd0\xcf\x11\xe0"
_XLRD_ERROR_CODES = {
    "#NULL!": 0x00,
    "#DIV/0!": 0x07,
    "#VALUE!": 0x0F,
    "#REF!": 0x17,
    "#NAME?": 0x1D,
    "#NUM!": 0x24,
    "#N/A": 0x2A,
}


def _uses_openpyxl(xlfile):
    if IRONPY:
        return False
    signature = _read_signature(xlfile)
    if signature.startswith(_ZIP_SIGNATURE):
        return True
    if signature.startswith(_OLE2_SIGNATURE):
        return False
    return not _is_legacy_workbook(xlfile)


def _read_signature(xlfile):
    try:
        with open(str(xlfile), "rb") as xlstream:
            return xlstream.read(len(_OLE2_SIGNATURE))
    except (IOError, OSError):
        return b""


def _is_legacy_workbook(xlfile):
    return op.splitext(str(xlfile))[1].lower() == ".xls"


def _legacy_xlrd():
    """Return xlrd, giving ``xlrd.book`` a ``time.clock`` where Python lacks it.

    xlrd predates the ``time.clock`` removal in Python 3.8. On such Pythons
    ``xlrd.book.time`` is replaced once, permanently, with a namespace that
    adds ``clock``; other modules' ``time`` is untouched.

    Note:
        The check-and-assign is idempotent and never undone, so concurrent
        readers cannot observe the shim being removed mid-read.
    """
    import time

    import xlrd
    import xlrd.book

    if not hasattr(xlrd.book.time, "clock"):
        import types

        xlrd.book.time = types.SimpleNamespace(clock=time.perf_counter, time=time.time)
    return xlrd


def _openpyxl_cell_value(cell, epoch):
    value = cell["value"]
    if cell["data_type"] == "e":
        return _XLRD_ERROR_CODES.get(value, value)
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(
        value, (datetime.datetime, datetime.date, datetime.time, datetime.timedelta)
    ):
        from openpyxl.utils.datetime import to_excel

        return float(to_excel(value, epoch=epoch))
    return value


def _openpyxl_sheet_rows(xlsheet):
    """Return the worksheet's rows as xlrd reports them.

    The sheet XML is parsed directly, so only cells stored in the file are
    visited and numbers keep their stored value even under a date format.
    Rows and columns span the cells holding a value and the merged ranges;
    cells that only carry formatting are ignored, as in xlrd.

    Note:
        ``xlsheet`` must come from a ``read_only`` workbook that is still
        open; its private ``_get_source``/``_shared_strings`` are read.
    """
    from openpyxl.utils.cell import range_boundaries
    from openpyxl.worksheet._reader import WorkSheetParser

    epoch = xlsheet.parent.epoch
    values = {}
    with xlsheet._get_source() as src:
        parser = WorkSheetParser(
            src,
            xlsheet._shared_strings,
            data_only=True,
            epoch=epoch,
            date_formats=set(),
            timedelta_formats=set(),
        )
        for _, row in parser.parse():
            for cell in row:
                if cell["value"] is not None:
                    values[(cell["row"], cell["column"])] = _openpyxl_cell_value(
                        cell, epoch
                    )
    nrows = max([r for r, _ in values] or [0])
    ncols = max([c for _, c in values] or [0])
    if parser.merged_cells is not None:
        for merged in parser.merged_cells.mergeCell:
            _, _, max_col, max_row = range_boundaries(merged.ref)
            nrows = max(nrows, max_row)
            ncols = max(ncols, max_col)
    return [
        [values.get((r, c), "") for c in range(1, ncols + 1)]
        for r in range(1, nrows + 1)
    ]


def _xlrd_sheet_rows(xlsheet):
    return [[cell.value for cell in row] for row in xlsheet.get_rows()]


def _open_sheets(xlfile, sheet_names=None):
    """Return ``(name, rows)`` for each worksheet, in workbook order.

    Picks the reader once per file. Rows are lists of xlrd-equivalent values
    whatever the reader; ``sheet_names``, when given, skips the other sheets
    before their rows are read.
    """
    if _uses_openpyxl(xlfile):
        with open(str(xlfile), "rb") as xlstream:
            xlwb = openpyxl.load_workbook(xlstream, read_only=True, data_only=True)
            try:
                return [
                    (sheet.title, _openpyxl_sheet_rows(sheet))
                    for sheet in xlwb.worksheets
                    if not sheet_names or sheet.title in sheet_names
                ]
            finally:
                xlwb.close()
    xlwb = _legacy_xlrd().open_workbook(str(xlfile))
    return [
        (sheet.name, _xlrd_sheet_rows(sheet))
        for sheet in xlwb.sheets()
        if not sheet_names or sheet.name in sheet_names
    ]


def _read_xlsheet(xlsheetrows, columns=None, datatype=None, headers=True):
    columns = columns or []
    xlsheetheader = xlsheetrows[0] if headers and xlsheetrows else []
    xlsheetdata = []
    for drow in xlsheetrows[1 if headers else 0 :]:
        if columns:
            drow = dict([x for x in zip(columns, drow)])
        elif datatype:
            drow = datatype(drow)
        xlsheetdata.append(drow)
    return {"headers": xlsheetheader, "rows": xlsheetdata}


def load(xlfile, sheets=None, columns=None, datatype=None, headers=True):
    """Read data from Excel file.

    Args:
        xlfile (str): full path of the excel file to read
        sheets (list, optional): worksheets to read. Defaults to all the sheets.
        columns (list, optional): Names to give to the columns.
            It builds a dictionary for each row with the column name and value.
            If none given (default), it returns a simple list of values.
        datatype (type, optional): Type of the data. Defaults to None.
        headers (bool, optional): Whether to use the first row as headers.
            Defaults to True.

    Returns:
        (dict[str, dict[str, Any]]): Excel data grouped by worksheet.
            Each worksheet is a dictionary with `headers` and `rows`.
    """
    xldata = {}
    for xlsheet_name, xlsheetrows in _open_sheets(xlfile, sheets):
        xldata[xlsheet_name] = _read_xlsheet(
            xlsheetrows,
            columns=columns,
            datatype=datatype,
            headers=headers,
        )
    return xldata


def _write_xlsxwriter(xlfile, datadict):
    xlwb = xlsxwriter.Workbook(xlfile)
    for xlsheetname, xlsheetdata in datadict.items():
        xlsheet = xlwb.add_worksheet(xlsheetname)
        for idx, data in enumerate(xlsheetdata):
            xlsheet.write_row(idx, 0, data)
    xlwb.close()


def _write_openpyxl(xlfile, datadict):
    xlwb = openpyxl.Workbook()
    if datadict:
        xlwb.remove(xlwb.active)
    for xlsheetname, xlsheetdata in datadict.items():
        xlsheet = xlwb.create_sheet(xlsheetname)
        for data in xlsheetdata:
            xlsheet.append(list(data))
    xlwb.save(xlfile)


def _writer_for(xlfile):
    if IRONPY:
        return _write_xlsxwriter
    if _is_legacy_workbook(xlfile):
        raise ValueError(
            "Cannot write legacy .xls workbooks; save as .xlsx: {}".format(xlfile)
        )
    return _write_openpyxl


def dump(xlfile, datadict):
    """Write data to Excel file.

    Creates a worksheet for each item of the input dictionary. An empty
    dictionary still produces a workbook with one empty worksheet.

    Args:
        xlfile (str): full path of the target excel file
        datadict (dict[str, list]): dictionary of worksheets names and data

    Raises:
        ValueError: on CPython, ``xlfile`` ends in ``.xls``; openpyxl writes
            ``.xlsx`` content only, which Excel would reject under a legacy
            name. IronPython keeps writing through xlsxwriter unchanged.
    """
    _writer_for(xlfile)(xlfile, datadict)
